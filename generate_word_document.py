"""
Script to generate complete Word Document (.docx) for all 15 DevOps Lab Tasks
Candidate: Gayatri Patil (PRN: 23102A0016, Roll: 16)
"""
import os
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_full_document():
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    COLOR_PRIMARY = RGBColor(14, 165, 233)     # Medical Sky Blue
    COLOR_DARK = RGBColor(30, 41, 59)          # Dark Slate
    COLOR_SECONDARY = RGBColor(99, 102, 241)   # Indigo
    HEX_HEADER_BG = "0EA5E9"
    HEX_ROW_ALT = "F8FAFC"
    HEX_CODE_BG = "0F172A"

    # ==========================================
    # 1. COVER PAGE
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_dept = title_p.add_run("DEPARTMENT OF COMPUTER ENGINEERING\nDEVOPS & CONTINUOUS TESTING LABORATORY\n")
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_DARK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(36)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_proj_lbl = p_proj.add_run("ACADEMIC LAB MANUAL & PROJECT DOSSIER\n")
    run_proj_lbl.font.size = Pt(12)
    run_proj_lbl.font.bold = True
    run_proj_lbl.font.color.rgb = COLOR_SECONDARY

    run_proj_title = p_proj.add_run("SELENIUM TESTING PIPELINE FOR A VACCINATION REMINDER PORTAL (VAXGUARD)\n")
    run_proj_title.font.size = Pt(22)
    run_proj_title.font.bold = True
    run_proj_title.font.color.rgb = COLOR_PRIMARY

    run_proj_sub = p_proj.add_run("End-to-End DevOps Automation: Git, Jenkins CI/CD, Docker Containerization & Ansible Provisioning\n")
    run_proj_sub.font.size = Pt(11)
    run_proj_sub.font.italic = True
    run_proj_sub.font.color.rgb = RGBColor(100, 116, 139)

    p_spacer2 = doc.add_paragraph()
    p_spacer2.paragraph_format.space_after = Pt(48)

    # Student Details Box
    student_table = doc.add_table(rows=5, cols=2)
    student_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    student_data = [
        ("Candidate Name:", "GAYATRI PATIL"),
        ("Student ID / PRN:", "23102A0016"),
        ("Roll Number:", "16"),
        ("Course / Semester:", "B.Tech / B.E. Computer Engineering (Sem VI)"),
        ("Academic Year:", "2026 - 2027")
    ]
    for i, (label, val) in enumerate(student_data):
        row = student_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(3.8)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_DARK
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(11)
        r1.font.bold = (i < 3)
        if i < 3:
            r1.font.color.rgb = COLOR_PRIMARY
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)

    doc.add_page_break()

    # ==========================================
    # 2. CERTIFICATE PAGE
    # ==========================================
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cert_title = p_cert_title.add_run("CERTIFICATE OF COMPLETION\n")
    r_cert_title.font.size = Pt(18)
    r_cert_title.font.bold = True
    r_cert_title.font.color.rgb = COLOR_DARK

    p_cert_body = doc.add_paragraph()
    p_cert_body.paragraph_format.line_spacing = 1.3
    p_cert_body.paragraph_format.space_after = Pt(24)
    r_cb = p_cert_body.add_run(
        "This is to certify that Miss GAYATRI PATIL, bearing Student ID / PRN 23102A0016 and Roll Number 16, "
        "has satisfactorily performed and completed all the prescribed 15 laboratory experiments and DevOps deliverables "
        "for the project entitled 'Selenium Testing Pipeline for a Vaccination Reminder Portal (VaxGuard)' "
        "as part of the academic curriculum for the course of DevOps Laboratory during the academic year 2026-2027."
    )
    r_cb.font.size = Pt(11.5)

    cert_table = doc.add_table(rows=2, cols=3)
    cert_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Date of Examination", "Internal Examiner Sign", "Head of Department"]
    for j, h in enumerate(headers):
        c = cert_table.rows[0].cells[j]
        c.width = Inches(2.2)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(100, 116, 139)

        c_space = cert_table.rows[1].cells[j]
        c_space.width = Inches(2.2)
        p_s = c_space.paragraphs[0]
        p_s.add_run("\n\n_______________________\nSignature")
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # ==========================================
    # 3. TABLE OF CONTENTS
    # ==========================================
    p_toc_head = doc.add_paragraph()
    r_th = p_toc_head.add_run("INDEX OF 15 DEVOPS LABORATORY ASSIGNMENTS")
    r_th.font.size = Pt(16)
    r_th.font.bold = True
    r_th.font.color.rgb = COLOR_PRIMARY

    index_table = doc.add_table(rows=16, cols=4)
    index_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(0.8), Inches(3.8), Inches(1.2), Inches(1.2)]
    headers_idx = ["Task #", "Deliverable Assignment Title", "Date", "Faculty Sign"]
    for j, h in enumerate(headers_idx):
        cell = index_table.rows[0].cells[j]
        cell.width = col_widths[j]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

    task_titles = [
        "Problem Definition, Stakeholder Analysis and Frozen Scope",
        "Agile Planning, User Stories, Product Backlog and DevOps Workflow",
        "Requirements, Multi-Tier Architecture and Spring Boot Tech Setup",
        "Git and GitHub Repository Initialization, .gitignore and Branch Policy",
        "Feature Development with Branching: National Immunization Schedule Engine",
        "MVP Completion, Simulated Merge Conflict Resolution and Release Tag v1.0.0",
        "Jenkins Installation, SCM Polling Trigger and Continuous Integration Job",
        "Pipeline as Code: Declarative 8-Stage Jenkinsfile with Environment Parameters",
        "Selenium Automated Test Design: 5 Critical User Journeys and Screenshot Listener",
        "Continuous Testing in Jenkins: Quality Gate Enforcement, Defect Fix & Rerun",
        "Docker Multi-Stage Image Build, Non-Root Hardening and Container Lifecycle",
        "Jenkins-Docker Continuous Deployment: Container Rollout and Health Gate",
        "Configuration Management Script: Server Prerequisites & Ansible Playbook",
        "Automated Provisioning, Idempotency Validation (changed=0) and Fast Rollback",
        "Final Release Release Dossier, Live Demonstration Guide and Complete Viva Prep"
    ]

    for i, title in enumerate(task_titles, start=1):
        row = index_table.rows[i]
        c0, c1, c2, c3 = row.cells
        c0.width, c1.width, c2.width, c3.width = col_widths

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run(f"Task {i:02d}").font.bold = True

        p1 = c1.paragraphs[0]
        p1.add_run(title)

        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run("Completed")

        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.add_run("Passed")

        bg = HEX_ROW_ALT if (i % 2 == 0) else "FFFFFF"
        for c in (c0, c1, c2, c3):
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)

    doc.add_page_break()

    # ==========================================
    # 4. LOAD AND APPEND ALL 15 TASKS
    # ==========================================
    task_files = sorted(glob.glob("tasks_for_teacher/Task_*.md"))
    for file_path in task_files:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        lines = content.split("\n")
        in_code_block = False
        code_buffer = []

        for line in lines:
            line_str = line.strip()

            # Skip redundant metadata headers in body
            if line_str.startswith("**Student Name:**") or line_str.startswith("**PRN / Student ID:**") or line_str.startswith("**Roll Number:**"):
                continue

            # Code Blocks
            if line_str.startswith("```"):
                if in_code_block:
                    in_code_block = False
                    # Emit code block
                    code_text = "\n".join(code_buffer)
                    cp = doc.add_paragraph()
                    cp.paragraph_format.left_indent = Inches(0.2)
                    cp.paragraph_format.space_before = Pt(4)
                    cp.paragraph_format.space_after = Pt(6)
                    r_code = cp.add_run(code_text)
                    r_code.font.name = "Consolas"
                    r_code.font.size = Pt(9)
                    r_code.font.color.rgb = RGBColor(16, 185, 129)
                    code_buffer = []
                else:
                    in_code_block = True
                continue

            if in_code_block:
                code_buffer.append(line)
                continue

            # Heading 1
            if line_str.startswith("# "):
                h1_text = line_str[2:].strip()
                h1 = doc.add_heading(level=1)
                h1.paragraph_format.space_before = Pt(16)
                h1.paragraph_format.space_after = Pt(6)
                r = h1.add_run(h1_text)
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
                continue

            # Heading 2
            if line_str.startswith("## "):
                h2_text = line_str[3:].strip()
                h2 = doc.add_heading(level=2)
                h2.paragraph_format.space_before = Pt(12)
                h2.paragraph_format.space_after = Pt(4)
                r = h2.add_run(h2_text)
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = COLOR_SECONDARY
                continue

            # Heading 3
            if line_str.startswith("### "):
                h3_text = line_str[4:].strip()
                h3 = doc.add_heading(level=3)
                h3.paragraph_format.space_before = Pt(8)
                h3.paragraph_format.space_after = Pt(2)
                r = h3.add_run(h3_text)
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = COLOR_DARK
                continue

            # Blockquotes / Speech
            if line_str.startswith(">"):
                quote_text = line_str.lstrip("> ").strip()
                qp = doc.add_paragraph()
                qp.paragraph_format.left_indent = Inches(0.3)
                qp.paragraph_format.space_before = Pt(4)
                qp.paragraph_format.space_after = Pt(4)
                r = qp.add_run(quote_text)
                r.font.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(15, 23, 42)
                continue

            # Bullet points
            if line_str.startswith("- ") or line_str.startswith("* "):
                bullet_text = line_str[2:].strip()
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(2)
                r = bp.add_run(bullet_text)
                r.font.size = Pt(10)
                continue

            # Standard paragraph
            if line_str:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(line_str)
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_DARK

        doc.add_page_break()

    # Output file
    output_filename = "DevOps_Lab_Manual_Gayatri_Patil_23102A0016.docx"
    doc.save(output_filename)
    print(f"[SUCCESS] Word Document generated: {output_filename}")

if __name__ == "__main__":
    create_full_document()
